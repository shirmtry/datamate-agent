import json
import re
import subprocess
import os
import requests
from pathlib import Path
from collections import Counter
import pandas as pd

# ================= CẤU HÌNH =================
BASE_DIR = Path(r"C:\Users\dinht\Downloads\train")
TRAIN_DIR = BASE_DIR / "train"
DATA_DIR = TRAIN_DIR / "data"
QUESTIONS_FILE = TRAIN_DIR / "questions.jsonl"
ANSWERS_FILE = TRAIN_DIR / "answers.jsonl"
LOG_FILE = BASE_DIR / "teacher_training_log.txt"

# Cấu hình Ollama
OLLAMA_URL = "http://localhost:11434/api/generate"
MODEL_NAME = "qwen2.5-coder:7b"
MAX_RETRIES = 3  

# ================= HÀM HỖ TRỢ =================
def compare_answer(pred, gold):
    if pred is None and gold is None: return True
    if pred is None or gold is None: return False
    pt, gt = pred.get("type"), gold.get("type")
    if pt != gt: return False
    p_ans, g_ans = pred.get("answer"), gold.get("answer")
    if pt == "unanswerable": return True
    if pt in ("number", "percent"):
        try: 
            return abs(float(str(p_ans).replace(',', '')) - float(str(g_ans).replace(',', ''))) <= (1e-6 if pt=="number" else 0.01)
        except: 
            return str(p_ans) == str(g_ans)
    if pt == "list":
        p = [str(x).strip().lower() for x in (p_ans or [])]
        g = [str(x).strip().lower() for x in (g_ans or [])]
        return p == g or Counter(p) == Counter(g)
    return str(p_ans).strip().lower() == str(g_ans).strip().lower()

def load_biz_data(biz_id):
    biz_dir = DATA_DIR / biz_id
    dfs = {}
    if not biz_dir.exists(): return dfs
    
    for f in biz_dir.iterdir():
        if f.name.startswith("~$") or not f.is_file(): continue
        try:
            if f.suffix.lower() == ".csv":
                try: dfs[f.stem] = pd.read_csv(f, encoding='utf-8')
                except: dfs[f.stem] = pd.read_csv(f, encoding='cp1258')
            elif f.suffix.lower() == ".xlsx":
                import openpyxl
                wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
                for sheet_name in wb.sheetnames:
                    sheet = wb[sheet_name]
                    data = list(sheet.values)
                    if not data: continue
                    
                    # Tìm dòng header
                    best_row, max_valid = 0, 0
                    for i in range(min(10, len(data))):
                        valid = sum(1 for v in data[i] if v is not None and str(v) != '')
                        if valid > max_valid: max_valid, best_row = valid, i
                    
                    headers = data[best_row]
                    rows = data[best_row+1:]
                    
                    clean_headers = []
                    for idx, h in enumerate(headers):
                        h_str = str(h).strip() if h is not None else f"Unnamed_{idx}"
                        clean_headers.append(h_str)
                    
                    # Xử lý lệch số cột
                    clean_rows = []
                    for r in rows:
                        r_list = list(r)
                        if len(r_list) > len(clean_headers): r_list = r_list[:len(clean_headers)]
                        elif len(r_list) < len(clean_headers): r_list.extend([None] * (len(clean_headers) - len(r_list)))
                        clean_rows.append(r_list)
                        
                    df_s = pd.DataFrame(clean_rows, columns=clean_headers)
                    df_s = df_s.loc[:, ~df_s.columns.astype(str).str.contains('^Unnamed')].dropna(how='all')
                    
                    key = f.stem if len(wb.sheetnames) == 1 else f"{f.stem}_{sheet_name}"
                    dfs[key] = df_s
                wb.close()
        except Exception as e:
            print(f"⚠️ Bỏ qua file {f.name} do lỗi: {e}")
    return dfs

def execute_ai_code(code_str, dfs, question, as_of_date):
    code_str = re.sub(r'^\s*(import|from)\s+.*$', '', code_str, flags=re.MULTILINE)
    local_vars = {"dfs": dfs, "pd": pd, "question": question, "as_of_date": as_of_date, "ans": None}
    try:
        exec(code_str, {"__builtins__": __builtins__}, local_vars)
        return local_vars.get("ans"), None
    except Exception as e:
        return None, f"Lỗi Runtime: {str(e)}"

def call_ollama(prompt):
    """Gọi trực tiếp Ollama API, bỏ qua DSH để tránh nhiễu prompt"""
    payload = {
        "model": MODEL_NAME,
        "prompt": prompt,
        "stream": False,
        "options": {"temperature": 0.0} # Nhiệt độ 0 để AI trả lời chính xác, bớt sáng tạo
    }
    try:
        res = requests.post(OLLAMA_URL, json=payload, timeout=300)
        if res.status_code == 200:
            response_text = res.json().get("response", "")
            # Trích xuất code
            match = re.search(r"```(?:python)?\n(.*?)```", response_text, re.DOTALL)
            if match:
                return match.group(1).strip(), None
            else:
                return None, f"Không tìm thấy code. AI nói: {response_text[:200]}..."
        else:
            return None, f"Lỗi HTTP: {res.status_code}"
    except Exception as e:
        return None, f"Lỗi kết nối Ollama: {str(e)}"

# ================= TIẾN TRÌNH CHÍNH =================
def main():
    if not QUESTIONS_FILE.exists() or not ANSWERS_FILE.exists():
        print("Không tìm thấy file questions.jsonl hoặc answers.jsonl")
        return

    questions = [json.loads(l) for l in open(QUESTIONS_FILE, "r", encoding="utf-8") if l.strip()]
    gold_answers = {json.loads(l)["id"]: json.loads(l) for l in open(ANSWERS_FILE, "r", encoding="utf-8") if l.strip()}
    
    questions = questions[:10] 
    correct_count = 0
    
    with open(LOG_FILE, "w", encoding="utf-8") as log:
        for i, q in enumerate(questions):
            q_id = q["id"]
            biz_id = q["business_id"]
            question = q["question"]
            as_of_date = q.get("as_of_date", "2026-12-31")
            gold = gold_answers.get(q_id)
            
            if not gold: continue
                
            print(f"\n[{i+1}/{len(questions)}] 🎯 Đang dạy AI câu: {q_id}")
            print(f"Câu hỏi: {question}")
            
            dfs = load_biz_data(biz_id)
            if not dfs:
                print("  ⚠️ Không có dữ liệu.")
                continue
                
            schema_info = "\n".join([f"- Bảng '{name}': {list(df.columns)}" for name, df in dfs.items()])
            log.write(f"\n{'='*60}\n[{q_id}] {question}\nGold: {gold}\n")
            
            previous_code = ""
            error_feedback = ""
            
            for attempt in range(MAX_RETRIES):
                print(f"  -> Lần thử {attempt + 1}/{MAX_RETRIES}...")
                result = None 
                
                # Prompt tối ưu, không bị DSH làm nhiễu
                prompt = f"""
Bạn là một Chuyên gia Phân tích Dữ liệu với Python và thư viện Pandas.
Nhiệm vụ của bạn là viết code Python để giải quyết câu hỏi sau:
Câu hỏi: "{question}"
Ngày chốt (as_of_date): {as_of_date}
Dữ liệu đã nạp sẵn trong dictionary `dfs`. Key là tên bảng/sheet.
Cấu trúc các bảng hiện có:
{schema_info}

QUY TẮC BẮT BUỘC:
1. TUYỆT ĐỐI KHÔNG DÙNG `pd.read_csv` hay `pd.read_excel`. Dữ liệu đã có sẵn trong `dfs`.
2. Kết quả cuối cùng phải được gán vào biến `ans`.
3. KHÔNG HARDCODE ĐÁP ÁN. Viết code tổng quát dựa trên logic thực tế. Chú ý xử lý ngày tháng và quy đổi đơn vị (thùng/lốc/lon) nếu có.
4. Code phải nằm trong block ```python ... ```.
"""
                if attempt > 0:
                    prompt += f"""
--------------------------------------------------
LẦN TRƯỚC BẠN ĐÃ VIẾT CODE NHƯ SAU VÀ BỊ SAI:
```python
{previous_code}
```

KẾT QUẢ CHẠY CODE ĐÓ: {result}
LỖI GẶP PHẢI: {error_feedback}
ĐÁP ÁN ĐÚNG CẦN ĐẠT LÀ: {gold['answer']}

Hãy phân tích lỗi sai và viết lại code CHÍNH XÁC HƠN.
"""
                
                code, dsh_err = call_ollama(prompt)
                
                if dsh_err:
                    print(f"  ❌ Lỗi AI: {dsh_err}")
                    error_feedback = dsh_err
                    continue
                    
                previous_code = code
                result, runtime_err = execute_ai_code(code, dfs, question, as_of_date)
                
                log.write(f"\n--- Lần {attempt+1} ---\nCode:\n{code}\n\nResult: {result}\nError: {runtime_err}\n")
                
                if runtime_err:
                    print(f"  ⚠️ Lỗi Code: {runtime_err}")
                    error_feedback = runtime_err
                    continue
                    
                pred = {"type": gold["type"], "answer": result}
                
                if compare_answer(pred, gold):
                    print(f"  ✅ THÀNH CÔNG! AI đã giải đúng.")
                    log.write(f"  ✅ KẾT QUẢ ĐÚNG!\n")
                    correct_count += 1
                    break
                else:
                    print(f"  ⚠️ SAI. AI ra: {result} | Gold: {gold['answer']}")
                    error_feedback = f"Kết quả ({result}) không khớp đáp án."
                    log.write(f"  ⚠️ SAI. AI ra: {result} | Gold: {gold['answer']}\n")

            if attempt == MAX_RETRIES - 1:
                print(f"  ❌ Thất bại sau {MAX_RETRIES} lần.")
                log.write(f"  ❌ THẤT BẠI.\n")

    print(f"\n🎉 HOÀN TẤT! Tỷ lệ đúng: {correct_count}/{len(questions)}")
    print(f"📄 Xem chi tiết tại: {LOG_FILE}")

if __name__ == "__main__":
    main()