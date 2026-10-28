import json
import re
import os
import sqlite3
import requests
import warnings
from pathlib import Path
import pandas as pd
import numpy as np
from datetime import datetime, timedelta

warnings.filterwarnings('ignore')

# ================= CẤU HÌNH =================
OLLAMA_URL = "http://localhost:11434/api/generate"
PLANNER_MODEL = "qwen2.5:3b"

# 🔴 ĐIỀN API KEY DEEPSEEK CỦA BẠN VÀO ĐÂY
DEEPSEEK_API_KEY = "sk-0b8f42c32f864fbd8cbb9412593bf54e" 
DEEPSEEK_URL = "https://api.deepseek.com/v1/chat/completions"
DEEPSEEK_MODEL = "deepseek-chat"

BASE_DIR = Path("C:/Users/dinht/Downloads/train")
TRAIN_DIR = BASE_DIR / "train"
DATA_DIR = TRAIN_DIR / "data"

def _detect_sep(fp):
    with open(fp, 'r', encoding='utf-8', errors='ignore') as f:
        line = f.readline()
        if ';' in line: return ';'
        if '\t' in line: return '\t'
        return ','

def load_biz_data(biz_id, data_dir):
    biz_dir = Path(data_dir) / biz_id
    dfs = {}
    if not biz_dir.exists(): return dfs
    for f in biz_dir.iterdir():
        if f.name.startswith("~$") or not f.is_file(): continue
        if f.suffix.lower() in [".csv", ".xlsx"]:
            try:
                if f.suffix.lower() == ".csv":
                    sep = _detect_sep(f)
                    for enc in ['utf-8', 'cp1258', 'latin1']:
                        try:
                            dfs[f.stem] = pd.read_csv(f, encoding=enc, sep=sep, on_bad_lines='skip', low_memory=False)
                            break
                        except: continue
                elif f.suffix.lower() == ".xlsx":
                    import openpyxl
                    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
                    for sn in wb.sheetnames:
                        data = list(wb[sn].values)
                        if not data: continue
                        best_row = max(range(min(10, len(data))), key=lambda i: sum(1 for v in data[i] if v is not None))
                        headers = [str(h).strip() if h is not None else f"Unnamed_{idx}" for idx, h in enumerate(data[best_row])]
                        rows = data[best_row+1:]
                        clean = [list(r)[:len(headers)] + [None]*(len(headers)-len(r)) for r in rows]
                        df_s = pd.DataFrame(clean, columns=headers)
                        key = f.stem if len(wb.sheetnames)==1 else f"{f.stem}_{sn}"
                        dfs[key] = df_s.dropna(how='all')
                    wb.close()
            except: pass
    return dfs

# ================= SCHEMA SIÊU CHI TIẾT =================
def get_detailed_schema(dfs):
    schema_str = ""
    for k, v in dfs.items():
        # Bỏ qua các cột Unnamed
        valid_cols = [c for c in v.columns if not str(c).startswith('Unnamed')]
        v = v[valid_cols]
        if len(valid_cols) == 0: continue
        
        schema_str += f"\n### Bảng: '{k}'\n"
        for col in valid_cols:
            dtype = str(v[col].dtype)
            if dtype == 'object' or dtype == 'string':
                # Lấy 5 giá trị unique cho cột chữ (Trạng thái, Tên hàng...)
                uniques = v[col].dropna().astype(str).unique()[:5]
                uniq_str = ", ".join([f"'{u}'" for u in uniques])
                schema_str += f"  - {col} (Chữ): {uniq_str}\n"
            else:
                # Cột số/ngày: Lấy 1 giá trị mẫu
                sample = str(v[col].dropna().iloc[0]) if not v[col].dropna().empty else "N/A"
                schema_str += f"  - {col} ({dtype}): Mẫu '{sample}'\n"
    return schema_str

def parse_date_safe(series):
    s = series.astype(str).str.strip()
    res = pd.to_datetime(s, errors='coerce', format='%Y-%m-%d')
    if res.notna().sum() > 0.5*len(res): return res
    res = pd.to_datetime(s, errors='coerce', dayfirst=True)
    if res.notna().sum() > 0.5 * len(res): return res
    return pd.to_datetime(s, errors='coerce')

def clean_money(s):
    return pd.to_numeric(s.astype(str).str.replace(r'[^0-9.-]', '', regex=True), errors='coerce').fillna(0)

def clean_number(s):
    return pd.to_numeric(s.astype(str).str.replace(r'[^0-9.-]', '', regex=True), errors='coerce').fillna(0)

def execute_ai_code(code_str, dfs, question, as_of_date):
    if not code_str: return None, None, None, "Không có code"
    
    def replace_read(match):
        file_str = match.group(1)
        file_name = Path(file_str).stem
        for k in dfs.keys():
            if file_name.lower() in k.lower() or k.lower() in file_name.lower():
                return f"dfs['{k}']"
        return f"(_ for _ in ()).throw(Exception('LỖI: File {file_str} không tồn tại. Các file hiện có: {list(dfs.keys())}'))"
        
    code_str = re.sub(r"pd\.read_csv\(['\"](.*?)['\"].*?\)", replace_read, code_str)
    code_str = re.sub(r"pd\.read_excel\(['\"](.*?)['\"].*?\)", replace_read, code_str)

    wrapper = f"""
def __solve__(dfs, question, as_of_date):
{chr(10).join('    ' + line for line in code_str.split(chr(10)))}
    return locals().get('ans', None), locals().get('reason', None), locals().get('ans_type', None)
"""
    env = {"__builtins__": __builtins__, "pd": pd, "np": np, "datetime": datetime, "timedelta": timedelta,
           "parse_date_safe": parse_date_safe, "clean_money": clean_money, "clean_number": clean_number}
    lv = {}
    try:
        exec(wrapper, env, lv)
        res = lv["__solve__"](dfs, question, as_of_date)
        return res[0], res[1], res[2], None
    except Exception as e:
        return None, None, None, f"Lỗi Runtime: {str(e)}"

def format_ans(ans, ans_type, question):
    if ans is None: return None
    if isinstance(ans, pd.DataFrame): return ans # Trả về DF để check_correct xử lý
    if isinstance(ans, (pd.Series, tuple)): ans = ans.tolist()
    if hasattr(ans, 'item'): ans = ans.item()
    if isinstance(ans, pd.Timestamp): return ans.strftime("%Y-%m-%d")
    if isinstance(ans, bool): return ans
    if isinstance(ans, (int, float, np.integer, np.floating)):
        v = float(ans)
        if v.is_integer(): v = int(v)
        if (ans_type == "percent") or ("%" in question) or ("tỷ lệ" in question.lower()):
            if abs(v) <= 1.0 and v != 0: v = round(v*100, 2)
            return round(v, 2)
        return v
    if isinstance(ans, list): return [str(x) for x in ans if pd.notna(x)]
    return str(ans).strip()

def check_correct(pred, gold):
    if pred is None or gold is None: return False
    g_ans = gold["answer"]
    g_type = gold.get("type", "")
    aliases = gold.get("alias", [])
    
    if g_type in ("number", "percent"):
        try: return abs(float(pred) - float(g_ans)) <= (1e-6 if g_type=="number" else 0.01)
        except: return False
        
    if g_type == "list":
        # Nếu pred là DataFrame, cố gắng lấy cột đầu tiên
        if isinstance(pred, pd.DataFrame):
            if pred.empty: return False
            pred = pred.iloc[:, 0].tolist()
            
        p = [str(x).strip().lower() for x in (pred or [])]
        g = [str(x).strip().lower() for x in g_ans]
        
        # Kiểm tra khớp chính xác
        if p == g: return True
        # Kiểm tra khớp alias
        for a in aliases:
            if isinstance(a, list) and p == [str(x).lower() for x in a]: return True
        return False
        
    p_s = str(pred).strip().lower()
    g_s = str(g_ans).strip().lower()
    if p_s == g_s: return True
    for a in aliases:
        if isinstance(a, str) and p_s == a.lower(): return True
    return False

# ================= AGENT DEEPSEEK =================
def _ollama(model, prompt, timeout=120, num_predict=100):
    try:
        r = requests.post(OLLAMA_URL, json={"model": model, "prompt": prompt, "stream": False, "options": {"temperature": 0.0, "num_predict": num_predict}}, timeout=timeout)
        return r.json().get("response", "")
    except Exception: return ""

def planner_plan(question, as_of, schema, readme):
    prompt = f"""Câu hỏi: "{question}"
Ngày chốt: {as_of}
Schema: {schema[:1500]}
Lập kế hoạch giải quyết trong 1-2 câu ngắn gọn (VD: Lấy bảng A, lọc ngày B, tính tổng cột C):"""
    return _ollama(PLANNER_MODEL, prompt, timeout=60, num_predict=80).strip()

def deepseek_coder(question, as_of_date, plan, schema, readme):
    prompt = f"""Viết code Python Pandas giải quyết câu hỏi.
Câu hỏi: "{question}"
Ngày chốt: {as_of_date}
Kế hoạch: {plan}
DỮ LIỆU THỰC TẾ (Dùng ĐÚNG tên cột này):
{schema}

QUY TẮC BẮT BUỘC (NẾU SAI SẼ BỊ LỖI):
1. "Hôm nay", "tháng này" phải dựa vào ngày {as_of_date}. KHÔNG dùng `datetime.now()`.
2. Gán kết quả vào `ans`. TUYỆT ĐỐI KHÔNG DÙNG `print()`.
3. KHÔNG import thư viện, KHÔNG dùng read_csv. Chỉ dùng `dfs['tên_bảng']`.
4. Cột số (tiền, số lượng) phải được làm sạch bằng `clean_money(df['cột'])` hoặc `clean_number(df['cột'])` TRƯỚC KHI TÍNH TOÁN.
5. Cột ngày dùng `parse_date_safe(df['cột'])`.
6. NẾU CÂU HỎI YÊU CẦU LIỆT KÊ/TOP (VD: "top 3 mặt hàng", "những món nào"): `ans` phải là một **LIST** các chuỗi tên (VD: `ans = ['A', 'B']`). TUYỆT ĐỐI KHÔNG trả về DataFrame.
7. Nếu câu hỏi hỏi "doanh thu", "tổng giá trị": tính tổng TIỀN. Nếu hỏi "bán được bao nhiêu": tính tổng SỐ LƯỢNG.
8. Nếu `df.empty` (không có dữ liệu thỏa mãn), gán `ans = None` và `reason = "Không có dữ liệu"`.
9. CHỈ TRẢ VỀ CODE TRONG ```python ... ```. KHÔNG giải thích.
"""
    headers = {"Authorization": f"Bearer {DEEPSEEK_API_KEY}", "Content-Type": "application/json"}
    payload = {"model": DEEPSEEK_MODEL, "messages": [{"role": "user", "content": prompt}], "temperature": 0.0, "max_tokens": 1000}
    try:
        res = requests.post(DEEPSEEK_URL, headers=headers, json=payload, timeout=120)
        if res.status_code == 200:
            text = res.json()["choices"][0]["message"]["content"]
            m = re.search(r"```(?:python)?\n(.*?)```", text, re.DOTALL)
            return m.group(1).strip() if m else None, None
        return None, f"Lỗi API {res.status_code}"
    except Exception as e: return None, str(e)

def deepseek_review(question, as_of_date, schema, code, output, gold_answer, gold_type):
    # Nếu output là DataFrame, chuyển thành string để đưa vào prompt
    output_str = output.to_string() if isinstance(output, pd.DataFrame) else str(output)
    
    prompt = f"""Code chạy sai hoặc ra kết quả không khớp.
Câu hỏi: {question}
Ngày chốt: {as_of_date}
Đáp án đúng (Loại: {gold_type}): {gold_answer}
Kết quả code chạy ra: {output_str[:500]}

Code lỗi:
```python
{code}
```
Hãy kiểm tra lỗi. Lưu ý:
- Cột số có dấu phẩy (VD: '42,000') phải dùng `clean_money()`.
- Nếu đáp án là List, hãy dùng `.tolist()` để trả về List, không trả về DataFrame.
- Kiểm tra lại xem đã dùng đúng ngày `as_of_date` chưa.

Viết lại code đã sửa. TUYỆT ĐỐI KHÔNG DÙNG `print()`, chỉ gán vào `ans`.
CHỈ TRẢ VỀ CODE TRONG ```python ... ```.
"""
    headers = {"Authorization": f"Bearer {DEEPSEEK_API_KEY}", "Content-Type": "application/json"}
    payload = {"model": DEEPSEEK_MODEL, "messages": [{"role": "user", "content": prompt}], "temperature": 0.0, "max_tokens": 1000}
    try:
        res = requests.post(DEEPSEEK_URL, headers=headers, json=payload, timeout=120)
        if res.status_code == 200:
            text = res.json()["choices"][0]["message"]["content"]
            m = re.search(r"```(?:python)?\n(.*?)```", text, re.DOTALL)
            return m.group(1).strip() if m else None
        return None
    except Exception: return None