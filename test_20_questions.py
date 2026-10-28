DEEPSEEK_API_KEY = "sk-0b8f42c32f864fbd8cbb9412593bf54e"
import json
import re
import os
import requests
import warnings
from pathlib import Path
import pandas as pd
import numpy as np
from datetime import datetime, timedelta

warnings.filterwarnings('ignore')

# ================= CẤU HÌNH =================
# 🔴 DÁN API KEY THẬT CỦA BẠN VÀO ĐÂY
DEEPSEEK_API_KEY = "XXX"

QUESTIONS_TO_RUN = 100 

BASE_DIR = Path(r"C:\Users\dinht\Downloads\train")
PUBLIC_DIR = BASE_DIR / "public_test"
QUESTIONS_FILE = PUBLIC_DIR / "questions.jsonl"
DATA_DIR = PUBLIC_DIR / "data"

OUTPUT_FILE = BASE_DIR / "v3.jsonl"
LOG_FILE = BASE_DIR / "logv3.txt"

DEEPSEEK_URL = "https://api.deepseek.com/v1/chat/completions"
MODEL_NAME = "deepseek-chat"
MAX_RETRIES = 3

# ================= LUẬT LỆ (RULEBOOK) =================
RULEBOOK_SUMMARY = """
BẠN PHẢI TUÂN THỦ CÁC LUẬT SAU:
1. DOANH THU: Chỉ tính đơn KHÔNG HỦY và KHÔNG TRẢ HÀNG. Đơn trả hàng (dù chỉ 1 phần) thì bỏ cả đơn. Đơn đang giao vẫn tính.
2. SỐ ĐƠN: Chỉ đếm đơn không hủy, không trả.
3. THỜI GIAN: "Tháng M/Y" là từ ngày 1 đến hết ngày cuối tháng. "N ngày gần nhất" tính cả hôm nay. "Tuần" bắt đầu từ Thứ 2.
4. XẾP HẠNG: Bán chạy = theo số lượng quy về đơn vị cơ bản. Lãi gộp = Doanh thu - Giá vốn.
5. CÔNG NỢ: Nợ = Tổng đơn - Đã trả (tính đến hết ngày X).
6. KHÔNG TRẢ LỜI ĐƯỢC: Nếu thiếu dữ liệu hoặc không tìm thấy, gán `ans = None` và `reason = "Giải thích lý do"`.
"""

# ================= HÀM HỖ TRỢ =================
def detect_separator(filepath):
    with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
        line = f.readline()
        if ';' in line: return ';'
        if '\t' in line: return '\t'
        return ','

def load_biz_data(biz_id):
    biz_dir = DATA_DIR / biz_id
    dfs = {}
    if not biz_dir.exists(): return dfs
    for f in biz_dir.iterdir():
        if f.name.startswith("~$") or not f.is_file(): continue
        try:
            if f.suffix.lower() == ".csv":
                sep = detect_separator(f)
                try: dfs[f.stem] = pd.read_csv(f, encoding='utf-8', sep=sep, on_bad_lines='skip', low_memory=False)
                except: dfs[f.stem] = pd.read_csv(f, encoding='cp1258', sep=sep, on_bad_lines='skip', low_memory=False)
            elif f.suffix.lower() == ".xlsx":
                import openpyxl
                wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
                for sheet_name in wb.sheetnames:
                    data = list(wb[sheet_name].values)
                    if not data: continue
                    best_row = max(range(min(10, len(data))), key=lambda i: sum(1 for v in data[i] if v is not None))
                    headers = [str(h).strip() if h is not None else f"Unnamed_{idx}" for idx, h in enumerate(data[best_row])]
                    rows = data[best_row+1:]
                    clean_rows = [list(r)[:len(headers)] + [None]*(len(headers)-len(r)) for r in rows]
                    df_s = pd.DataFrame(clean_rows, columns=headers)
                    dfs[f.stem if len(wb.sheetnames)==1 else f"{f.stem}_{sheet_name}"] = df_s.dropna(how='all')
                wb.close()
        except Exception as e: 
            pass
    return dfs

def parse_date_safe(series):
    s = series.astype(str).str.strip()
    res = pd.to_datetime(s, errors='coerce', format='%Y-%m-%d')
    if res.notna().sum() > 0.5 * len(res): return res
    res = pd.to_datetime(s, errors='coerce', dayfirst=True)
    if res.notna().sum() > 0.5 * len(res): return res
    return pd.to_datetime(s, errors='coerce')

def clean_money(series):
    return pd.to_numeric(series.astype(str).str.replace(r'[^0-9.-]', '', regex=True), errors='coerce').fillna(0)

def clean_number(series):
    return pd.to_numeric(series.astype(str).str.replace(r'[^0-9.-]', '', regex=True), errors='coerce').fillna(0)

def execute_ai_code(code_str, dfs, question, as_of_date):
    # Sửa lỗi Scope: Đưa dfs vào thẳng global_env thay vì truyền qua biến phụ
    wrapper_code = f"""
def __solve__():
{chr(10).join('    ' + line for line in code_str.split(chr(10)))}
    return locals().get('ans', None), locals().get('reason', None), locals().get('ans_type', None)
"""
    global_env = {
        "__builtins__": __builtins__, 
        "pd": pd, "np": np, "datetime": datetime, "timedelta": timedelta,
        "parse_date_safe": parse_date_safe, "clean_money": clean_money, "clean_number": clean_number,
        "dfs": dfs, "question": question, "as_of_date": as_of_date # Đã đưa dfs vào đây
    }
    local_vars = {}
    try:
        exec(wrapper_code, global_env, local_vars)
        return local_vars["__solve__"]() + (None,)
    except Exception as e:
        return None, None, None, f"Lỗi Runtime: {str(e)}"

def call_deepseek(prompt):
    headers = {"Authorization": f"Bearer {DEEPSEEK_API_KEY}", "Content-Type": "application/json"}
    payload = {
        "model": MODEL_NAME,
        "messages": [
            {"role": "system", "content": "Bạn là chuyên gia phân tích dữ liệu Python/Pandas. Chỉ trả về code trong block ```python ... ```."},
            {"role": "user", "content": prompt}
        ],
        "temperature": 0.0
    }
    try:
        res = requests.post(DEEPSEEK_URL, headers=headers, json=payload, timeout=120)
        if res.status_code == 200:
            response_text = res.json()["choices"][0]["message"]["content"]
            match = re.search(r"```(?:python)?\n(.*?)```", response_text, re.DOTALL)
            if match: return match.group(1).strip(), None
            return None, "Không tìm thấy code."
        elif res.status_code == 401:
            return None, "LỖI 401: API KEY KHÔNG HỢP LỆ."
        else:
            return None, f"Lỗi API {res.status_code}: {res.text}"
    except Exception as e:
        return None, f"Lỗi kết nối: {str(e)}"

def format_output(q_id, ans, reason, ans_type, question):
    if ans is None:
        if reason: return {"id": q_id, "type": "unanswerable", "reason": str(reason)}
        return None 

    if isinstance(ans, pd.DataFrame):
        clean_df = ans.astype(object).where(pd.notnull(ans), None)
        return {"id": q_id, "type": "table", "answer": {"columns": [str(c) for c in clean_df.columns], "rows": clean_df.values.tolist()}}
    
    if isinstance(ans, pd.Series): ans = ans.tolist()
    if hasattr(ans, 'item'): ans = ans.item()

    if isinstance(ans, (pd.Timestamp, datetime)): 
        if "tháng" in question.lower() and "ngày" not in question.lower():
            return {"id": q_id, "type": "date", "answer": ans.strftime("%Y-%m")}
        return {"id": q_id, "type": "date", "answer": ans.strftime("%Y-%m-%d")}

    if isinstance(ans, (list, tuple)):
        clean_list = []
        for x in ans:
            if pd.isna(x): continue
            if isinstance(x, (int, float, np.integer, np.floating)): clean_list.append(float(x) if not float(x).is_integer() else int(x))
            elif isinstance(x, dict): clean_list.append(str(list(x.values())[0]))
            else: clean_list.append(str(x))
        return {"id": q_id, "type": "list", "answer": clean_list}

    ans_str = str(ans).strip()
    
    if "có" in question.lower() or "phải không" in question.lower():
        if ans_str.lower() in ["không", "sai", "false"]: return {"id": q_id, "type": "boolean", "answer": False}
        if ans_str.lower() in ["có", "đúng", "true"]: return {"id": q_id, "type": "boolean", "answer": True}

    if isinstance(ans, str) and ("phần trăm" in question.lower() or "%" in question or "tăng" in question.lower() or "giảm" in question.lower()):
        nums = re.findall(r'-?\d+(?:\.\d+)?', ans.replace(',', ''))
        if nums:
            val = float(nums[0])
            if "giảm" in ans.lower() or "kém" in ans.lower(): val = -abs(val)
            elif "tăng" in ans.lower() or "hơn" in ans.lower(): val = abs(val)
            return {"id": q_id, "type": "percent", "answer": round(val, 2)}

    if isinstance(ans, (int, float, np.integer, np.floating)):
        ans_float = float(ans)
        if ans_float.is_integer(): ans_float = int(ans_float)
        is_percent = (ans_type == "percent") or ("phần trăm" in question.lower()) or ("%" in question) or ("tỷ lệ" in question.lower())
        if is_percent:
            if abs(ans_float) <= 1.0 and ans_float != 0 and ("tăng" in question.lower() or "giảm" in question.lower() or "thay đổi" in question.lower()):
                ans_float = round(ans_float * 100, 2)
            return {"id": q_id, "type": "percent", "answer": round(ans_float, 2)}
        return {"id": q_id, "type": "number", "answer": ans_float}

    if isinstance(ans, bool): return {"id": q_id, "type": "boolean", "answer": ans}

    ans_str = re.sub(r'^["\']|["\']$', '', ans_str)
    return {"id": q_id, "type": "string", "answer": ans_str}

# ================= CHẠY CHÍNH =================
def main():
    all_questions = [json.loads(l) for l in open(QUESTIONS_FILE, "r", encoding="utf-8") if l.strip()]
    questions = all_questions[:QUESTIONS_TO_RUN]
    print(f"🚀 Bắt đầu chạy {len(questions)} câu...")
    
    processed = set()
    if OUTPUT_FILE.exists():
        with open(OUTPUT_FILE, "r", encoding="utf-8") as f:
            for line in f:
                if line.strip(): processed.add(json.loads(line)["id"])
        if processed:
            print(f"♻️ Đã tìm thấy {len(processed)} câu trong file kết quả cũ, sẽ bỏ qua...")

    with open(OUTPUT_FILE, "a", encoding="utf-8") as f_out, open(LOG_FILE, "a", encoding="utf-8") as log:
        for i, q in enumerate(questions):
            q_id = q["id"]
            if q_id in processed: continue
            
            print(f"\n[{i+1}/{len(questions)}] 🎯 {q_id} | {q['business_id']}")
            dfs = load_biz_data(q["business_id"])
            
            if not dfs:
                print("  ⚠️ Không có dữ liệu -> ABSTAIN.")
                f_out.write(json.dumps({"id": q_id, "type": "abstain"}, ensure_ascii=False) + "\n")
                f_out.flush()
                continue
                
            schema_info = "\n".join([f"- Bảng '{name}': {list(df.columns)}" for name, df in dfs.items()])
            prev_code, err_feedback = "", ""
            final_rec = None
            
            for attempt in range(MAX_RETRIES):
                print(f"  -> Lần thử {attempt + 1}/{MAX_RETRIES}...")
                prompt = f"""
{RULEBOOK_SUMMARY}

Câu hỏi: "{q['question']}"
Ngày chốt (as_of_date): {q.get('as_of_date')}
Dữ liệu đã nạp sẵn trong dict `dfs`. Key là tên bảng:
{schema_info}

QUY TẮC CODE BẮT BUỘC (NẾU VI PHẠM SẼ BỊ LỖI NGAY LẬP TỨC):
1. KHÔNG CẦN import thư viện. `pd`, `np`, `datetime`, `parse_date_safe`, `clean_money`, `clean_number`, `dfs` đã có sẵn.
2. KHÔNG DÙNG pd.read_csv/read_excel. Dữ liệu đã có sẵn trong `dfs`.
3. Gán kết quả cuối cùng vào biến `ans`.
4. Code trong block ```python ... ```.
5. TUYỆT ĐỐI KHÔNG GIẢ ĐỊNH có sẵn các biến lạ (như `qd_map`, `he_so_map`, `dm_match`, `name_to_sku`...). Mọi biến PHẢI ĐƯỢC ĐỊNH NGHĨA ngay trong code.
6. LÀM SẠCH DỮ LIỆU SỐ: Nếu cột chứa chữ, hãy dùng các hàm có sẵn: `clean_money(df['Cột'])` hoặc `clean_number(df['Cột'])`. KHÔNG TỰ VIẾT regex làm sạch.
7. XỬ LÝ NGÀY THÁNG: Dùng hàm `parse_date_safe(df['tên_cột'])`.
8. XỬ LÝ TÊN: Dùng `df['Cột'].str.contains('tên_cần_tìm', case=False, na=False)` thay vì dùng `==`.
9. NẾU LÀ LIST: Nếu câu hỏi yêu cầu liệt kê top, hãy trả về một List các chuỗi tên (VD: `ans = ["A", "B"]`). KHÔNG trả về List các Dictionary.
10. Nếu không tìm thấy dữ liệu, gán `ans = None` và `reason = "Lý do cụ thể"`.
"""
                if attempt > 0:
                    prompt += f"\nLẦN TRƯỚC BỊ LỖI:\n```python\n{prev_code}\n```\nLỖI GẶP PHẢI: {err_feedback}\nHãy sửa lại code chính xác hơn!"

                code, api_err = call_deepseek(prompt)
                
                if api_err:
                    print(f"  ❌ {api_err}")
                    if "401" in api_err: return
                    err_feedback = api_err
                    continue
                    
                prev_code = code
                result, reason, ans_type, run_err = execute_ai_code(code, dfs, q["question"], q.get("as_of_date"))
                
                log.write(f"\n{'='*40}\n{q_id}\nCode:\n{code}\nRes: {result}\nReason: {reason}\nErr: {run_err}\n")
                
                if run_err:
                    print(f"  ⚠️ Lỗi Code: {run_err}")
                    err_feedback = run_err
                    continue
                    
                final_rec = format_output(q_id, result, reason, ans_type, q["question"])
                if final_rec:
                    print(f"  ✅ Đã ra kết quả: {result}")
                    break
                else:
                    print("  ⚠️ Không có kết quả trả về (None).")
                    err_feedback = "Code chạy không lỗi nhưng không gán gì vào biến `ans`. Hãy kiểm tra lại logic."
            
            if final_rec is not None:
                f_out.write(json.dumps(final_rec, ensure_ascii=False) + "\n")
            else:
                print("  ❌ Thất bại sau 3 lần thử -> Ghi ABSTAIN (0 điểm) để tránh bị trừ điểm.")
                f_out.write(json.dumps({"id": q_id, "type": "abstain"}, ensure_ascii=False) + "\n")
            
            f_out.flush()

    print(f"\n🎉 Đã chạy xong. File kết quả: {OUTPUT_FILE}")

if __name__ == "__main__":
    main()