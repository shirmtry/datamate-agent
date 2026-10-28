### 📁 File 4: `run_pipeline.py` (Nhạc trưởng)
# File này sẽ gọi 3 con AI trên, chạy code, nếu lỗi thì gọi Tester sửa, nếu sai quá 3 lần thì đánh `abstain` (0 điểm).
import json
import re
import os
import warnings
from pathlib import Path
import pandas as pd
import numpy as np
from datetime import datetime, timedelta

# Import 3 Agents
from agent_planner import create_plan
from agent_coder import write_code
from agent_tester import fix_code

warnings.filterwarnings('ignore')

# ================= CẤU HÌNH =================
QUESTIONS_TO_RUN = 100 
BASE_DIR = Path(r"C:\Users\dinht\Downloads\train")
PUBLIC_DIR = BASE_DIR / "public_test"
QUESTIONS_FILE = PUBLIC_DIR / "questions.jsonl"
DATA_DIR = PUBLIC_DIR / "data"

OUTPUT_FILE = BASE_DIR / "submission_local.jsonl"

# ================= HÀM HỖ TRỢ =================
def detect_separator(filepath):
    with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
        line = f.readline()
        if ';' in line: return ';'
        if '\t' in line: return '\t'
        return ','

def load_biz_data(biz_id):
    biz_dir = DATA_DIR / biz_id
    dfs = {}
    if not biz_dir.exists(): return dfs
    for f in biz_dir.iterdir():
        if f.name.startswith("~$") or not f.is_file(): continue
        try:
            if f.suffix.lower() == ".csv":
                sep = detect_separator(f)
                try: dfs[f.stem] = pd.read_csv(f, encoding='utf-8', sep=sep, on_bad_lines='skip', low_memory=False)
                except: dfs[f.stem] = pd.read_csv(f, encoding='cp1258', sep=sep, on_bad_lines='skip', low_memory=False)
            elif f.suffix.lower() == ".xlsx":
                import openpyxl
                wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
                for sheet_name in wb.sheetnames:
                    data = list(wb[sheet_name].values)
                    if not data: continue
                    best_row = max(range(min(10, len(data))), key=lambda i: sum(1 for v in data[i] if v is not None))
                    headers = [str(h).strip() if h is not None else f"Unnamed_{idx}" for idx, h in enumerate(data[best_row])]
                    rows = data[best_row+1:]
                    clean_rows = [list(r)[:len(headers)] + [None]*(len(headers)-len(r)) for r in rows]
                    df_s = pd.DataFrame(clean_rows, columns=headers)
                    dfs[f.stem if len(wb.sheetnames)==1 else f"{f.stem}_{sheet_name}"] = df_s.dropna(how='all')
                wb.close()
        except Exception as e: 
            pass
    return dfs

def parse_date_safe(series):
    s = series.astype(str).str.strip()
    res = pd.to_datetime(s, errors='coerce', format='%Y-%m-%d')
    if res.notna().sum() > 0.5 * len(res): return res
    res = pd.to_datetime(s, errors='coerce', dayfirst=True)
    if res.notna().sum() > 0.5 * len(res): return res
    return pd.to_datetime(s, errors='coerce')

def clean_money(series):
    return pd.to_numeric(series.astype(str).str.replace(r'[^0-9.-]', '', regex=True), errors='coerce').fillna(0)

def clean_number(series):
    return pd.to_numeric(series.astype(str).str.replace(r'[^0-9.-]', '', regex=True), errors='coerce').fillna(0)

def execute_ai_code(code_str, dfs, question, as_of_date):
    if not code_str: return None, None, None, "Không có code"
    wrapper_code = f"""
def __solve__():
{chr(10).join('    ' + line for line in code_str.split(chr(10)))}
    return locals().get('ans', None), locals().get('reason', None), locals().get('ans_type', None)
"""
    global_env = {
        "__builtins__": __builtins__, 
        "pd": pd, "np": np, "datetime": datetime, "timedelta": timedelta,
        "parse_date_safe": parse_date_safe, "clean_money": clean_money, "clean_number": clean_number,
        "dfs": dfs, "question": question, "as_of_date": as_of_date
    }
    local_vars = {}
    try:
        exec(wrapper_code, global_env, local_vars)
        return local_vars["__solve__"]() + (None,)
    except Exception as e:
        return None, None, None, f"Lỗi Runtime: {str(e)}"

def format_output(q_id, ans, reason, ans_type, question):
    if ans is None:
        if reason: return {"id": q_id, "type": "unanswerable", "reason": str(reason)}
        return None 

    if isinstance(ans, pd.DataFrame):
        clean_df = ans.astype(object).where(pd.notnull(ans), None)
        return {"id": q_id, "type": "table", "answer": {"columns": [str(c) for c in clean_df.columns], "rows": clean_df.values.tolist()}}
    if isinstance(ans, pd.Series): ans = ans.tolist()
    if hasattr(ans, 'item'): ans = ans.item()

    if isinstance(ans, (pd.Timestamp, datetime)): 
        if "tháng" in question.lower() and "ngày" not in question.lower():
            return {"id": q_id, "type": "date", "answer": ans.strftime("%Y-%m")}
        return {"id": q_id, "type": "date", "answer": ans.strftime("%Y-%m-%d")}

    if isinstance(ans, (list, tuple)):
        clean_list = []
        for x in ans:
            if pd.isna(x): continue
            if isinstance(x, (int, float, np.integer, np.floating)): clean_list.append(float(x) if not float(x).is_integer() else int(x))
            elif isinstance(x, dict): clean_list.append(str(list(x.values())[0]))
            else: clean_list.append(str(x))
        return {"id": q_id, "type": "list", "answer": clean_list}

    ans_str = str(ans).strip()
    if "có" in question.lower() or "phải không" in question.lower():
        if ans_str.lower() in ["không", "sai", "false"]: return {"id": q_id, "type": "boolean", "answer": False}
        if ans_str.lower() in ["có", "đúng", "true"]: return {"id": q_id, "type": "boolean", "answer": True}

    if isinstance(ans, str) and ("phần trăm" in question.lower() or "%" in question or "tăng" in question.lower() or "giảm" in question.lower()):
        nums = re.findall(r'-?\d+(?:\.\d+)?', ans.replace(',', ''))
        if nums:
            val = float(nums[0])
            if "giảm" in ans.lower() or "kém" in ans.lower(): val = -abs(val)
            elif "tăng" in ans.lower() or "hơn" in ans.lower(): val = abs(val)
            return {"id": q_id, "type": "percent", "answer": round(val, 2)}

    if isinstance(ans, (int, float, np.integer, np.floating)):
        ans_float = float(ans)
        if ans_float.is_integer(): ans_float = int(ans_float)
        is_percent = (ans_type == "percent") or ("phần trăm" in question.lower()) or ("%" in question) or ("tỷ lệ" in question.lower())
        if is_percent:
            if abs(ans_float) <= 1.0 and ans_float != 0 and ("tăng" in question.lower() or "giảm" in question.lower() or "thay đổi" in question.lower()):
                ans_float = round(ans_float * 100, 2)
            return {"id": q_id, "type": "percent", "answer": round(ans_float, 2)}
        return {"id": q_id, "type": "number", "answer": ans_float}

    if isinstance(ans, bool): return {"id": q_id, "type": "boolean", "answer": ans}
    ans_str = re.sub(r'^["\']|["\']$', '', ans_str)
    return {"id": q_id, "type": "string", "answer": ans_str}

# ================= CHẠY CHÍNH =================
def main():
    all_questions = [json.loads(l) for l in open(QUESTIONS_FILE, "r", encoding="utf-8") if l.strip()]
    questions = all_questions[:QUESTIONS_TO_RUN]
    print(f"🚀 Bắt đầu chạy Pipeline Local Model trên {len(questions)} câu...")
    
    processed = set()
    if OUTPUT_FILE.exists():
        with open(OUTPUT_FILE, "r", encoding="utf-8") as f:
            for line in f:
                if line.strip(): processed.add(json.loads(line)["id"])
    
    with open(OUTPUT_FILE, "a", encoding="utf-8") as f_out:
        for i, q in enumerate(questions):
            q_id = q["id"]
            if q_id in processed: continue
            
            print(f"\n[{i+1}/{len(questions)}] 🎯 {q_id} | {q['business_id']}")
            dfs = load_biz_data(q["business_id"])
            
            if not dfs:
                print("  ⚠️ Không có dữ liệu -> ABSTAIN.")
                f_out.write(json.dumps({"id": q_id, "type": "abstain"}, ensure_ascii=False) + "\n")
                f_out.flush()
                continue
                
            schema_info = "\n".join([f"- Bảng '{name}': {list(df.columns)}" for name, df in dfs.items()])
            
            # 1. PLANNER (3B) LÊN KẾ HOẠCH
            print("  -> Planner (3B) đang lên kế hoạch...")
            plan = create_plan(q["question"], q.get("as_of_date"), schema_info)
            
            # 2. CODER (1.5B) VIẾT CODE
            print("  -> Coder (1.5B) đang viết code...")
            code = write_code(q["question"], plan, schema_info)
            
            final_rec = None
            
            # 3. VÒNG LẶP CHẠY THỬ & SỬA LỖI (7B)
            for attempt in range(3):
                print(f"     * Vòng thử {attempt+1}/3...")
                if not code:
                    print("       ❌ Coder không viết được code.")
                    break
                    
                result, reason, ans_type, err = execute_ai_code(code, dfs, q["question"], q.get("as_of_date"))
                
                if err:
                    print(f"       ⚠️ Lỗi Code: {err}")
                    print("       -> Tester (7B) đang sửa lỗi...")
                    code = fix_code(q["question"], code, err)
                    continue
                    
                final_rec = format_output(q_id, result, reason, ans_type, q["question"])
                if final_rec:
                    print(f"       ✅ Chạy thành công! Kết quả: {result}")
                    break
                else:
                    print("       ⚠️ Code chạy nhưng không có kết quả (None).")
                    code = fix_code(q["question"], code, "Code chạy không lỗi nhưng không gán gì vào biến `ans`.")
            
            # Xử lý lưu file
            if final_rec is not None:
                f_out.write(json.dumps(final_rec, ensure_ascii=False) + "\n")
            else:
                print("  ❌ Thất bại sau 3 lần thử -> ABSTAIN (0 điểm).")
                f_out.write(json.dumps({"id": q_id, "type": "abstain"}, ensure_ascii=False) + "\n")
            
            f_out.flush()

    print(f"\n🎉 Đã chạy xong Pipeline Local. File kết quả: {OUTPUT_FILE}")

if __name__ == "__main__":
    main()