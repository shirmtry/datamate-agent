"""
utils.py — Helper chống lỗi dữ liệu bẩn cho DataMate (bản nâng cấp).

Điểm mạnh:
  - Đọc csv (tự dò sep/encoding), xlsx (tự dò header), sqlite.
  - parse_date_safe xử lý "1 thg 1 2025", dd/mm/yyyy, yyyy-mm-dd... từng phần tử.
  - clean_money/clean_number hiểu "1.234.567 đ", "50k", "1,5tr", "2 triệu".
  - execute_ai_code: cô lập trong __solve__, thay pd.read_* bằng get_df_by_keyword,
    pd.concat(df) an toàn, bắt mọi Exception.
"""
import io
import re
import sqlite3
import tokenize
import unicodedata
import warnings
from datetime import datetime, timedelta
from pathlib import Path

import numpy as np
import pandas as pd

warnings.filterwarnings("ignore")


# ================= NẠP DỮ LIỆU =================
def load_biz_data(biz_id, data_dir):
    biz_dir = Path(data_dir) / str(biz_id)
    dfs = {}
    if not biz_dir.exists():
        return dfs
    for f in sorted(biz_dir.iterdir()):
        if f.name.startswith("~$") or not f.is_file():
            continue
        try:
            if f.suffix.lower() == ".csv":
                dfs[f.stem] = _read_csv(f)
            elif f.suffix.lower() == ".xlsx":
                dfs.update(_read_xlsx(f))
            elif f.suffix.lower() in (".sqlite", ".db"):
                dfs.update(_read_sqlite(f))
        except Exception:
            continue
    return dfs


def _read_csv(path):
    sample = path.read_bytes()[:8192].decode("utf-8", errors="ignore")
    sep = max([",", ";", "\t", "|"], key=lambda s: sample.count(s))
    for enc in ("utf-8-sig", "utf-8", "cp1258", "latin1"):
        try:
            df = pd.read_csv(path, encoding=enc, sep=sep, on_bad_lines="skip", low_memory=False)
            if df.shape[1] == 1:
                alt = pd.read_csv(path, encoding=enc, sep=None, engine="python", on_bad_lines="skip")
                if alt.shape[1] > 1:
                    df = alt
            return df
        except Exception:
            continue
    return pd.read_csv(path, encoding="utf-8", errors="ignore", on_bad_lines="skip")


def _read_xlsx(path):
    import openpyxl
    out = {}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for sn in wb.sheetnames:
            data = list(wb[sn].values)
            if not data:
                continue
            best = max(range(min(10, len(data))),
                       key=lambda i: sum(1 for v in data[i] if v is not None))
            headers = [str(h).strip() if h is not None else f"Unnamed_{i}"
                       for i, h in enumerate(data[best])]
            rows = [list(r)[:len(headers)] + [None] * (len(headers) - len(r))
                    for r in data[best + 1:]]
            key = f"{path.stem}_{sn}" if len(wb.sheetnames) > 1 else path.stem
            out[key] = pd.DataFrame(rows, columns=headers).dropna(how="all")
    finally:
        wb.close()
    return {k: v for k, v in out.items() if len(v) > 0}


def _read_sqlite(path):
    out = {}
    conn = sqlite3.connect(path)
    try:
        for (name,) in conn.execute("SELECT name FROM sqlite_master WHERE type='table'"):
            out[name] = pd.read_sql(f'SELECT * FROM "{name}"', conn)
    finally:
        conn.close()
    return out


def get_readme(biz_id, data_dir, max_chars=6000):
    p = Path(data_dir) / str(biz_id) / "README.md"
    if not p.exists():
        return ""
    text = p.read_text(encoding="utf-8", errors="ignore").strip()
    if len(text) <= max_chars:
        return text
    head = int(max_chars * 0.6)
    return text[:head] + "\n... (lược bớt) ...\n" + text[-(max_chars - head):]


# ================= SCHEMA =================
def get_full_schema(dfs):
    lines = []
    for name, df in dfs.items():
        cols = [str(c) for c in df.columns if not str(c).startswith("Unnamed")]
        if cols:
            lines.append(f"- BẢNG '{name}': {cols}")
    return "\n".join(lines)


def compact_schema(dfs):
    groups = {}
    for name, df in dfs.items():
        cols = tuple(str(c) for c in df.columns if not str(c).startswith("Unnamed"))
        if cols:
            groups.setdefault(cols, []).append(str(name))
    return "\n".join(f"- {' | '.join(names)}: {list(cols)}" for cols, names in groups.items())


# ================= TÌM BẢNG =================
def normalize_text(text):
    if pd.isna(text):
        return ""
    t = unicodedata.normalize("NFD", str(text).lower()).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", t)


def get_df_by_keyword(dfs, keywords):
    if isinstance(keywords, str):
        keywords = [keywords]
    kws = [k for k in (normalize_text(x) for x in keywords) if k]
    for kw in kws:
        for name in dfs:
            if normalize_text(name) == kw:
                return dfs[name]
    for kw in kws:
        cands = [n for n in dfs if kw in normalize_text(n)]
        if cands:
            pref = [n for n in cands if normalize_text(n).startswith(kw)]
            return dfs[min(pref or cands, key=lambda n: len(normalize_text(n)))]
    return None


def get_all_dfs_by_keyword(dfs, keywords):
    if isinstance(keywords, str):
        keywords = [keywords]
    kws = [k for k in (normalize_text(x) for x in keywords) if k]
    picked, seen = [], set()
    for kw in kws:
        for name, df in dfs.items():
            if kw in normalize_text(name) and df is not None and len(df) > 0:
                if id(df) not in seen:
                    seen.add(id(df))
                    picked.append(df)
    if not picked:
        return None
    return picked[0] if len(picked) == 1 else pd.concat(picked, ignore_index=True, sort=False)


# ================= LÀM SẠCH =================
def clean_money(series):
    """'1.234.567 đ' -> 1234567 ; '50k' -> 50000 ; '1,5tr' -> 1500000 ; '2 triệu' -> 2000000."""
    if series is None:
        return None
    scalar = not isinstance(series, pd.Series)
    s = pd.Series([series]) if scalar else series
    s = s.astype(str).str.lower().str.strip()

    mult = pd.Series(1.0, index=s.index)
    mult[s.str.contains(r"\d\s*k\b|\d\s*nghìn", regex=True, na=False)] = 1e3
    mult[s.str.contains(r"\d\s*tr\b|\d\s*triệu|\d\s*tr\.", regex=True, na=False)] = 1e6
    mult[s.str.contains(r"\d\s*tỷ", regex=True, na=False)] = 1e9

    num = s.str.replace(r"[^0-9,.\-]", "", regex=True)
    num = num.str.replace(r"\.(?=\d{3}(\D|$))", "", regex=True)   # 1.234.567 -> 1234567
    num = num.str.replace(",", ".", regex=False)
    val = pd.to_numeric(num, errors="coerce").fillna(0) * mult
    return val.iloc[0] if scalar else val


def clean_number(series):
    return clean_money(series)


def parse_date_safe(series):
    """Parse từng phần tử: '1 thg 1 2025', '30/04/2025', '2025-04-30', '30-04-2025'..."""
    if series is None:
        return None
    scalar = not isinstance(series, pd.Series)
    s = pd.Series([series]) if scalar else series
    s = s.astype(str).str.strip()

    def _thg(x):
        m = re.search(r"(\d{1,2})\s*thg\s*(\d{1,2})\s*(\d{4})", str(x).lower())
        if m:
            d, mo, y = map(int, m.groups())
            try:
                return pd.Timestamp(year=y, month=mo, day=d)
            except ValueError:
                return pd.NaT
        return pd.NaT

    thg = s.apply(_thg)
    rest = thg.isna()
    if rest.any():
        part = s[rest]
        parsed = pd.to_datetime(part, errors="coerce", dayfirst=True)
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d", "%d.%m.%Y",
                    "%Y-%m-%d %H:%M:%S", "%d/%m/%Y %H:%M"):
            miss = parsed.isna()
            if not miss.any():
                break
            parsed = parsed.mask(miss, pd.to_datetime(part, errors="coerce", format=fmt))
        thg = pd.concat([thg[~rest], parsed]).reindex(s.index)

    return thg.iloc[0] if scalar else thg


def normalize_id(series):
    if series is None:
        return None
    if not isinstance(series, pd.Series):
        series = pd.Series([series])
    return series.astype(str).str.strip().str.upper()


def strip_comments(code):
    code = code or ""
    lines = code.splitlines()
    spans = {}
    try:
        for tok in tokenize.generate_tokens(io.StringIO(code).readline):
            if tok.type == tokenize.COMMENT:
                r, c0 = tok.start
                _, c1 = tok.end
                spans.setdefault(r - 1, []).append((c0, c1))
    except Exception:
        return "\n".join(l for l in lines if l.strip())
    for r, sp in spans.items():
        line = lines[r]
        for c0, c1 in sorted(sp, reverse=True):
            line = line[:c0] + line[c1:]
        lines[r] = line.rstrip()
    return "\n".join(l for l in lines if l.strip())


def clean_code_text(code):
    out = []
    for line in strip_comments(code).splitlines():
        s = line.strip()
        if re.match(r"^(print\s*\(|import\s+pandas|import\s+numpy|from\s+pandas\b|from\s+numpy\b)", s):
            continue
        if re.match(r"^as_of_date\s*=\s*(pd\.Timestamp|pd\.to_datetime|datetime\.strptime)", s):
            continue
        out.append(line)
    return "\n".join(out).strip()


# ================= THỰC THI CODE =================
_READ_CALL = re.compile(
    r"pd\.read_(?:csv|excel|table|sql)\s*\(\s*[rbf]?['\"]([^'\"]+)['\"][^)]*\)", re.IGNORECASE)
_BANNED = re.compile(r"\bopen\s*\(|\bsubprocess\b|\burllib\b|\bsocket\b")


def _replace_file_reads(code):
    def _sub(m):
        return f"get_df_by_keyword(dfs, [{Path(m.group(1)).stem!r}])"
    return _READ_CALL.sub(_sub, code)


class _PandasProxy:
    """pd.concat(df) với 1 DataFrame vẫn chạy."""
    def __init__(self, mod):
        object.__setattr__(self, "_mod", mod)

    def __getattr__(self, name):
        return getattr(self._mod, name)

    def concat(self, objs, *args, **kwargs):
        if isinstance(objs, pd.DataFrame):
            objs = [objs]
        return self._mod.concat(objs, *args, **kwargs)


def execute_ai_code(code_str, dfs, question=None, as_of_date=None):
    """Trả (ans, reason, ans_type, error). Không bao giờ raise."""
    if not code_str:
        return None, None, None, "Không có code"

    code = _replace_file_reads(code_str)
    if _BANNED.search(code):
        return None, None, None, "Lỗi: không được dùng open()/subprocess/network"

    if isinstance(as_of_date, str):
        try:
            as_of_date = pd.Timestamp(as_of_date)
        except Exception:
            pass

    body = "\n".join("    " + ln for ln in code.splitlines())
    wrapper = ("def __solve__(dfs, question, as_of_date):\n"
               f"{body}\n"
               "    return locals().get('ans'), locals().get('reason'), locals().get('ans_type')\n")

    env = {
        "__builtins__": __builtins__,
        "pd": _PandasProxy(pd), "np": np, "datetime": datetime, "timedelta": timedelta,
        "parse_date_safe": parse_date_safe, "clean_money": clean_money,
        "clean_number": clean_number, "normalize_id": normalize_id,
        "get_df_by_keyword": get_df_by_keyword,
        "get_all_dfs_by_keyword": get_all_dfs_by_keyword,
        "dfs": dfs, "question": question, "as_of_date": as_of_date,
    }
    local = {}
    try:
        exec(compile(wrapper, "<ai_code>", "exec"), env, local)
        ans, reason, ans_type = local["__solve__"](dfs, question, as_of_date)
        return ans, reason, ans_type, None
    except Exception as e:
        return None, None, None, f"Lỗi Runtime: {type(e).__name__}: {e}"
